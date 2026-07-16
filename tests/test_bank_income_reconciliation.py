from __future__ import annotations

from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from lector_facturas.api.app import create_app
from lector_facturas.api.store import ReviewStore
from lector_facturas.bank_income_reconciliation import (
    _match_manual_income,
    _parse_caixabank_xlsx,
)
from lector_facturas.bank_income_review_workbook import (
    build_manual_income_review_workbook,
    collect_manual_income_review,
)


class BankIncomeReconciliationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = TemporaryDirectory()
        self.finance_root = Path(self.tmp.name) / "finance"
        statement_dir = self.finance_root / "Artesta Store, S.L" / "2026" / "202603" / "statements" / "bank"
        statement_dir.mkdir(parents=True, exist_ok=True)
        self.statement_path = statement_dir / "Caixabank_movimientoscuenta_202603.xlsx"
        self._write_statement(self.statement_path)

        storage_path = Path(self.tmp.name) / "review_items.json"
        self.store = ReviewStore(storage_path=storage_path, finance_root=self.finance_root)
        self.store.upsert_shopify_payout_transactions([
            {
                "source_record_id": "shopify-1",
                "transaction_date": "2026-03-04T00:00:00Z",
                "payout_date": "2026-03-05T00:00:00Z",
                "payout_id": "po-mar-1",
                "net": "100.00",
                "currency": "EUR",
                "company_code": "SL",
            },
            {
                "source_record_id": "shopify-2",
                "transaction_date": "2026-02-27T00:00:00Z",
                "payout_date": "2026-02-28T00:00:00Z",
                "payout_id": "po-feb-1",
                "net": "80.00",
                "currency": "EUR",
                "company_code": "SL",
            },
            {
                "source_record_id": "shopify-3",
                "transaction_date": "2026-03-18T00:00:00Z",
                "payout_date": "2026-03-19T00:00:00Z",
                "payout_id": "po-mar-missing",
                "net": "125.00",
                "currency": "EUR",
                "company_code": "SL",
            },
        ])
        self.store.set_bank_manual_orders([
            {"company_code": "SL", "order_name": "MAN-OLD-1", "period_yyyymm": "202601", "order_date": "2026-03-18", "total": "150.00", "tags": ""},
            {"company_code": "SL", "order_name": "MAN-GRP-1", "period_yyyymm": "202603", "order_date": "2026-03-16", "total": "40.00", "tags": ""},
            {"company_code": "SL", "order_name": "MAN-GRP-2", "period_yyyymm": "202603", "order_date": "2026-03-20", "total": "50.00", "tags": ""},
            {"company_code": "SL", "order_name": "MAN-PARTIAL", "period_yyyymm": "202603", "order_date": "2026-03-24", "total": "120.00", "tags": ""},
        ])

        app = create_app()
        app.dependency_overrides.clear()
        from lector_facturas.api.app import get_store

        app.dependency_overrides[get_store] = lambda: self.store
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_run_sl_caixabank_income_reconciliation(self) -> None:
        response = self.client.post("/bank/income/sl/run", json={"period_yyyymm": "202603"})
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["company_code"], "SL")
        self.assertEqual(payload["files_processed"], 1)
        self.assertEqual(payload["transactions_imported"], 6)

        transactions = self.client.get("/bank/income/transactions?company_code=SL&period_yyyymm=202603&include_excluded=true").json()
        self.assertEqual(len(transactions), 6)

        by_amount = {item["amount"]: item for item in transactions}
        self.assertEqual(by_amount["100.00"]["classification"], "shopify_payout_candidate")
        self.assertEqual(by_amount["100.00"]["status"], "validated")
        self.assertEqual(by_amount["80.00"]["reason_code"], "shopify_outside_period")
        self.assertEqual(by_amount["60.00"]["status"], "ignored")
        self.assertEqual(by_amount["60.00"]["exclusion_reason"], "paypal_out_of_scope")
        self.assertEqual(by_amount["150.00"]["match_type"], "manual_order")
        self.assertEqual(by_amount["90.00"]["status"], "suggested")
        self.assertEqual(by_amount["90.00"]["match_type"], "manual_order_group")
        self.assertEqual(by_amount["30.00"]["reason_code"], "manual_partial_payment")

        reviews = self.client.get("/bank/income/reviews?company_code=SL&period_yyyymm=202603").json()
        reason_codes = {item["reason_code"] for item in reviews}
        self.assertIn("shopify_missing_bank_entry", reason_codes)
        self.assertIn("manual_order_unpaid", reason_codes)
        self.assertIn("manual_group_match_suggested", reason_codes)

    def test_manual_override_endpoint_updates_transaction(self) -> None:
        self.client.post("/bank/income/sl/run", json={"period_yyyymm": "202603"})
        transactions = self.client.get("/bank/income/transactions?company_code=SL&period_yyyymm=202603&include_excluded=true").json()
        group_tx = next(item for item in transactions if item["amount"] == "90.00")

        response = self.client.post(
            "/bank/income/matches/manual",
            json={
                "bank_transaction_id": group_tx["id"],
                "status": "validated",
                "reason_code": "manual_confirmed",
                "match_type": "manual_order_group",
                "target_ids": ["MAN-GRP-1", "MAN-GRP-2"],
                "notes": "Confirmed by operator",
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["status"], "validated")
        self.assertEqual(payload["source"], "manual")
        self.assertEqual(payload["target_ids"], ["MAN-GRP-1", "MAN-GRP-2"])

    def test_manual_match_prioritizes_recent_period_before_old_backlog(self) -> None:
        statement_dir = self.finance_root / "Artesta Store, S.L" / "2026" / "202601" / "statements" / "bank"
        statement_dir.mkdir(parents=True, exist_ok=True)
        path = statement_dir / "Caixabank_prioritized.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Movimientos_cuenta_0129605"
        ws.append(["Movimientos de la cuenta", "", "", "", "", ""])
        ws.append(["Importes expresados en euros", "", "", "", "", ""])
        ws.append(["Fecha", "Fecha valor", "Movimiento", "Más datos", "Importe", "Saldo"])
        ws.append(["24/01/2026", "26/01/2026", "TRANSF. A SU FAVOR", "00491803-ALAMO FORO SLU", 516.46, 36473.09])
        wb.save(path)

        self.store.set_bank_manual_orders([
            {"company_code": "SL", "order_name": "OLD-1", "period_yyyymm": "202511", "order_date": "2025-11-17", "total": "516.46", "tags": ""},
            {"company_code": "SL", "order_name": "OLD-2", "period_yyyymm": "202512", "order_date": "2025-12-16", "total": "516.46", "tags": ""},
            {"company_code": "SL", "order_name": "CUR-1", "period_yyyymm": "202601", "order_date": "2026-01-14", "total": "516.46", "tags": ""},
            {"company_code": "SL", "order_name": "CUR-2", "period_yyyymm": "202601", "order_date": "2026-01-20", "total": "516.46", "tags": ""},
        ])

        response = self.client.post("/bank/income/sl/run", json={"period_yyyymm": "202601"})
        self.assertEqual(response.status_code, 200)
        transactions = self.client.get("/bank/income/transactions?company_code=SL&period_yyyymm=202601&include_excluded=true").json()
        tx = next(item for item in transactions if item["amount"] == "516.46")

        self.assertEqual(tx["status"], "validated")
        self.assertEqual(tx["target_ids"], ["CUR-2"])

    def test_manual_match_ignores_candidates_outside_ten_day_window(self) -> None:
        tx = {
            "id": "tx-prepaid",
            "company_code": "SL",
            "period_yyyymm": "202512",
            "booking_date": "2025-12-24",
            "value_date": "2025-12-24",
            "amount": "516.46",
            "concept": "TRANSF. A SU FAVOR",
            "detail": "00491803-ALAMO FORO SLU",
            "classification": "manual_transfer_candidate",
            "is_excluded": False,
            "exclusion_reason": "",
        }
        match, _review, consumed = _match_manual_income(
            tx=tx,
            open_orders=[
                {"company_code": "SL", "order_name": "OLD-DEC", "period_yyyymm": "202512", "order_date": "2025-12-10", "total": "516.46", "tags": ""},
                {"company_code": "SL", "order_name": "FUTURE-JAN", "period_yyyymm": "202601", "order_date": "2026-01-14", "total": "516.46", "tags": ""},
            ],
            used_manual_order_ids=set(),
            current_period="202512",
        )

        self.assertEqual(match["status"], "pending_review")
        self.assertEqual(match["reason_code"], "manual_income_unmatched")
        self.assertEqual(match["target_ids"], [])
        self.assertEqual(consumed, [])

    def test_manual_match_picks_closest_candidate_within_ten_day_window(self) -> None:
        tx = {
            "id": "tx-window",
            "company_code": "SL",
            "period_yyyymm": "202603",
            "booking_date": "2026-03-10",
            "value_date": "2026-03-10",
            "amount": "516.46",
            "concept": "TRANSF. A SU FAVOR",
            "detail": "00491803-ALAMO FORO SLU",
            "classification": "manual_transfer_candidate",
            "is_excluded": False,
            "exclusion_reason": "",
        }
        match, _review, consumed = _match_manual_income(
            tx=tx,
            open_orders=[
                {"company_code": "SL", "order_name": "TOO-EARLY", "period_yyyymm": "202602", "order_date": "2026-02-20", "total": "516.46", "tags": ""},
                {"company_code": "SL", "order_name": "VALID-CLOSE", "period_yyyymm": "202603", "order_date": "2026-03-08", "total": "516.46", "tags": ""},
                {"company_code": "SL", "order_name": "VALID-FARTHER", "period_yyyymm": "202603", "order_date": "2026-03-15", "total": "516.46", "tags": ""},
            ],
            used_manual_order_ids=set(),
            current_period="202603",
        )

        self.assertEqual(match["status"], "validated")
        self.assertEqual(match["target_ids"], ["VALID-CLOSE"])
        self.assertEqual(consumed, ["VALID-CLOSE"])

    def test_parse_caixabank_xlsx_with_header_row_layout(self) -> None:
        path = Path(self.tmp.name) / "real_layout.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Movimientos_cuenta_0129605"
        ws.append(["Movimientos de la cuenta", "", "", "", "", ""])
        ws.append(["Importes expresados en euros", "", "", "", "", ""])
        ws.append(["Fecha", "Fecha valor", "Movimiento", "Más datos", "Importe", "Saldo"])
        ws.append(["09/01/2026", "09/01/2026", "TRANSF. A SU FAVOR", "00493609-TOP FLOOR SL", 2248.88, 74403.72])
        ws.append(["09/01/2026", "09/01/2026", "TRANSFER. EN DIV.", "CITINL2X   -STRIPE", 8986.29, 83390.01])
        wb.save(path)

        rows = _parse_caixabank_xlsx(path)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["booking_date"], "2026-01-09")
        self.assertEqual(rows[0]["concept"], "TRANSF. A SU FAVOR")
        self.assertEqual(rows[0]["detail"], "00493609-TOP FLOOR SL")
        self.assertEqual(str(rows[0]["amount"]), "2248.88")
        self.assertEqual(rows[1]["concept"], "TRANSFER. EN DIV.")
        self.assertEqual(rows[1]["detail"], "CITINL2X   -STRIPE")

    def test_build_manual_income_review_workbook(self) -> None:
        path = Path(self.tmp.name) / "manual_review.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Movimientos_cuenta_0129605"
        ws.append(["Movimientos de la cuenta", "", "", "", "", ""])
        ws.append(["Importes expresados en euros", "", "", "", "", ""])
        ws.append(["Fecha", "Fecha valor", "Movimiento", "Más datos", "Importe", "Saldo"])
        ws.append(["19/03/2026", "19/03/2026", "TRANSF. A SU FAVOR", "01280029-MISMAR GREEN SL", 150.00, 1390.00])
        ws.append(["20/03/2026", "20/03/2026", "TRANSF. A SU FAVOR", "01280029-MISMAR GREEN SL", 90.00, 1480.00])
        wb.save(path)

        self.store.set_bank_manual_orders([
            {"company_code": "SL", "order_name": "MAN-1", "period_yyyymm": "202603", "order_date": "2026-03-19", "total": "150.00", "tags": ""},
            {"company_code": "SL", "order_name": "MAN-2", "period_yyyymm": "202603", "order_date": "2026-03-20", "total": "120.00", "tags": ""},
        ])
        report = collect_manual_income_review(
            self.store,
            company_code="SL",
            statement_path=path,
            from_date="2026-03-01",
            to_date="2026-03-31",
        )

        workbook_bytes = build_manual_income_review_workbook(report)
        out_path = Path(self.tmp.name) / "out.xlsx"
        out_path.write_bytes(workbook_bytes)
        loaded = load_workbook(out_path)

        self.assertEqual(loaded.sheetnames, ["Summary", "Validated", "Conflicts", "Pending", "All Rows"])
        validated = loaded["Validated"]
        pending = loaded["Pending"]
        self.assertEqual(validated["A2"].value, "2026-03-19")
        self.assertEqual(validated["I2"].value, "MAN-1")
        self.assertEqual(pending["A2"].value, "2026-03-20")

    def _write_statement(self, path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Movimientos_cuenta_0129605"
        ws.append(["Movimientos de la cuenta", "", "", "", "", ""])
        ws.append(["Importes expresados en euros", "", "", "", "", ""])
        ws.append(["TRANSFER. EN DIV.", "2026-03-05", "2026-03-05", "CITINL2X   -STRIPE", 100.00, 1100.00])
        ws.append(["TRANSFER. EN DIV.", "2026-03-02", "2026-03-02", "CITINL2X   -STRIPE", 80.00, 1180.00])
        ws.append(["TRANSFER. EN DIV.", "2026-03-09", "2026-03-09", "PPLXLUL2   -PayPal Europe S.a.r.", 60.00, 1240.00])
        ws.append(["TRANSF. A SU FAVOR", "2026-03-19", "2026-03-19", "01280029-MISMAR GREEN SL", 150.00, 1390.00])
        ws.append(["TRANSF. A SU FAVOR", "2026-03-21", "2026-03-21", "00491803-ALAMO FORO SLU", 90.00, 1480.00])
        ws.append(["TRANSF. A SU FAVOR", "2026-03-22", "2026-03-22", "00496692-NULA STUDIO SL", 30.00, 1510.00])
        wb.save(path)
