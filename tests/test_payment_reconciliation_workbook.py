from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from lector_facturas.payment_reconciliation import ReconciliationReport
from lector_facturas.payment_reconciliation_workbook import build_reconciliation_workbook


def test_summary_includes_gift_cards_column() -> None:
    report = ReconciliationReport(
        period_yyyymm="202606",
        company_code="SL",
        gift_card_accounting_total=Decimal("43.80"),
        gift_card_payment_total=Decimal("43.80"),
    )

    workbook_bytes = build_reconciliation_workbook(report)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
    ws = workbook["Summary"]

    assert ws["E5"].value == "Gift Cards"
    assert ws["E6"].value == 43.8
    assert ws["E7"].value == 43.8
    assert ws["G6"].value == "=C6+D6+E6+F6"
    assert ws["G7"].value == "=C7+D7+E7+F7"
