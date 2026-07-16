from __future__ import annotations

import csv
import hashlib
import itertools
import re
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None  # type: ignore[assignment]

from lector_facturas.review_workflow import company_folder_name


_TWO_PLACES = Decimal("0.01")
_ZERO = Decimal("0.00")
_SHOPIFY_PATTERN = re.compile(r"(STRIPE|CITINL2X)", re.IGNORECASE)
_PAYPAL_PATTERN = re.compile(r"(PAYPAL|PPLXLUL2)", re.IGNORECASE)
_MANUAL_PATTERN = re.compile(r"TRANSF\.?\s+A\s+SU\s+FAVOR", re.IGNORECASE)
_MAX_MANUAL_PAYMENT_DELTA_DAYS = 10
_CAIXABANK_HEADER_ALIASES = {
    "booking_date": {"fecha"},
    "value_date": {"fecha valor"},
    "concept": {"movimiento"},
    "detail": {"mas datos"},
    "amount": {"importe"},
    "balance": {"saldo"},
}


@dataclass(frozen=True)
class BankIncomeRunResult:
    company_code: str
    period_yyyymm: str
    statement_dir: str
    files_processed: int
    transactions_imported: int
    validated_count: int
    suggested_count: int
    pending_review_count: int
    conflict_count: int
    ignored_count: int
    review_count: int


def run_sl_caixabank_income_reconciliation(store, *, period_yyyymm: str) -> BankIncomeRunResult:
    if not store.finance_root:
        raise RuntimeError("FINANCE_ROOT is required for bank income reconciliation.")

    statement_dir = _statement_dir(root=store.finance_root, company_code="SL", period_yyyymm=period_yyyymm)
    if not statement_dir.exists():
        raise FileNotFoundError(f"Statement directory not found: {statement_dir}")

    caixabank_files = sorted(
        path for path in statement_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in {".xls", ".xlsx", ".csv"}
        and "caixabank" in path.name.lower()
    )
    if not caixabank_files:
        raise FileNotFoundError(f"No Caixabank statements found in {statement_dir}")

    transactions: list[dict[str, Any]] = []
    for path in caixabank_files:
        file_hash = _file_hash(path)
        statement_file = {
            "id": _stable_id("bank-statement", "SL", period_yyyymm, str(path), file_hash),
            "company_code": "SL",
            "period_yyyymm": period_yyyymm,
            "bank_name": "Caixabank",
            "account_label": "Caixabank",
            "source_path": str(path),
            "file_name": path.name,
            "file_hash": file_hash,
            "parser_name": f"caixabank_{path.suffix.lower().lstrip('.')}",
            "imported_at": datetime.now(tz=UTC).isoformat().replace("+00:00", "Z"),
        }
        store.upsert_bank_statement_file(statement_file)
        for row_index, row in enumerate(_parse_caixabank_statement(path), start=1):
            fingerprint = _transaction_fingerprint(row)
            classification, is_excluded, exclusion_reason = _classify_transaction(row)
            transactions.append(
                {
                    "id": _stable_id("bank-income-tx", "SL", "Caixabank", fingerprint),
                    "statement_file_id": statement_file["id"],
                    "company_code": "SL",
                    "period_yyyymm": period_yyyymm,
                    "bank_name": "Caixabank",
                    "account_label": "Caixabank",
                    "booking_date": row["booking_date"],
                    "value_date": row["value_date"],
                    "amount": _format_decimal(row["amount"]),
                    "balance": _format_decimal(row["balance"]) if row["balance"] is not None else None,
                    "currency": "EUR",
                    "direction": "credit" if row["amount"] > 0 else "debit",
                    "concept": row["concept"],
                    "detail": row["detail"],
                    "classification": classification,
                    "is_excluded": is_excluded,
                    "exclusion_reason": exclusion_reason,
                    "fingerprint": fingerprint,
                    "source_row_number": row_index,
                    "raw_payload": {
                        "concept": row["concept"],
                        "detail": row["detail"],
                        "booking_date": row["booking_date"],
                        "value_date": row["value_date"],
                        "amount": _format_decimal(row["amount"]),
                        "balance": _format_decimal(row["balance"]) if row["balance"] is not None else None,
                    },
                }
            )

    store.replace_bank_transactions(company_code="SL", period_yyyymm=period_yyyymm, transactions=transactions)

    current_transactions = store.list_bank_income_transactions(
        company_code="SL",
        period_yyyymm=period_yyyymm,
        include_excluded=True,
    )
    manual_overrides = {
        item["bank_transaction_id"]: item
        for item in store.list_bank_income_matches(company_code="SL", period_yyyymm=period_yyyymm, source="manual")
    }

    previous_period = _previous_period(period_yyyymm)
    payouts = store.list_shopify_payout_aggregates(company_code="SL", from_period=previous_period, to_period=period_yyyymm)
    open_manual_orders = store.list_manual_open_orders(company_code="SL", period_yyyymm=period_yyyymm)

    used_payout_ids = {
        target_id
        for match in manual_overrides.values()
        if match.get("match_type") == "shopify_payout"
        for target_id in match.get("target_ids", [])
    }
    used_manual_order_ids = {
        target_id
        for target_id in store.list_consumed_manual_order_ids(company_code="SL")
    }
    used_manual_order_ids.update(
        target_id
        for match in manual_overrides.values()
        if str(match.get("match_type", "")).startswith("manual_order")
        for target_id in match.get("target_ids", [])
    )

    auto_matches: list[dict[str, Any]] = []
    reviews: list[dict[str, Any]] = []
    review_consumed_orders: set[str] = set(used_manual_order_ids)

    positive_transactions = sorted(
        (tx for tx in current_transactions if Decimal(str(tx["amount"])) > _ZERO),
        key=lambda item: (str(item.get("booking_date") or ""), str(item["id"])),
    )

    for tx in positive_transactions:
        if tx["id"] in manual_overrides:
            preserved = manual_overrides[tx["id"]]
            if str(preserved.get("match_type", "")).startswith("manual_order"):
                review_consumed_orders.update(preserved.get("target_ids", []))
            continue

        classification = str(tx.get("classification") or "")
        if classification == "paypal_income_ignored":
            auto_matches.append(
                _match_record(
                    tx,
                    status="ignored",
                    reason_code="paypal_income_ignored",
                    match_type="",
                    target_ids=[],
                    confidence=1,
                    is_excluded=True,
                    exclusion_reason="paypal_out_of_scope",
                )
            )
            continue
        if classification == "unknown_income":
            auto_matches.append(
                _match_record(
                    tx,
                    status="pending_review",
                    reason_code="unknown_income",
                    match_type="",
                    target_ids=[],
                    confidence=0,
                )
            )
            reviews.append(
                _review_record(
                    company_code="SL",
                    period_yyyymm=period_yyyymm,
                    bank_transaction_id=tx["id"],
                    review_type="transaction",
                    reason_code="unknown_income",
                    notes=f"Unclassified income: {tx.get('concept', '')} | {tx.get('detail', '')}".strip(" |"),
                )
            )
            continue
        if classification == "shopify_payout_candidate":
            match, review, payout_id = _match_shopify_payout(
                tx=tx,
                payouts=payouts,
                used_payout_ids=used_payout_ids,
                current_period=period_yyyymm,
            )
            auto_matches.append(match)
            if review:
                reviews.append(review)
            if payout_id:
                used_payout_ids.add(payout_id)
            continue
        if classification == "manual_transfer_candidate":
            match, review, consumed_orders = _match_manual_income(
                tx=tx,
                open_orders=open_manual_orders,
                used_manual_order_ids=used_manual_order_ids,
                current_period=period_yyyymm,
            )
            auto_matches.append(match)
            if review:
                reviews.append(review)
            used_manual_order_ids.update(consumed_orders)
            if match["status"] in {"validated", "suggested"}:
                review_consumed_orders.update(consumed_orders)
            continue

    matched_or_preserved_payouts = used_payout_ids | {
        target_id
        for match in auto_matches
        if match.get("match_type") == "shopify_payout"
        for target_id in match.get("target_ids", [])
    }
    for payout in payouts:
        if payout["target_id"] in matched_or_preserved_payouts:
            continue
        reviews.append(
            _review_record(
                company_code="SL",
                period_yyyymm=period_yyyymm,
                bank_transaction_id="",
                review_type="missing_expected",
                reason_code="shopify_missing_bank_entry",
                target_ids=[payout["target_id"]],
                notes=f"Missing Caixabank entry for Shopify payout {payout['target_id']} ({payout['net_amount']})",
            )
        )

    for order in open_manual_orders:
        if order["order_name"] in review_consumed_orders:
            continue
        reviews.append(
            _review_record(
                company_code="SL",
                period_yyyymm=period_yyyymm,
                bank_transaction_id="",
                review_type="missing_expected",
                reason_code="manual_order_unpaid",
                target_ids=[order["order_name"]],
                notes=f"Manual order still unpaid: {order['order_name']} ({order['total']})",
            )
        )

    store.replace_auto_bank_income_matches(company_code="SL", period_yyyymm=period_yyyymm, matches=auto_matches)
    store.replace_auto_bank_income_reviews(company_code="SL", period_yyyymm=period_yyyymm, reviews=reviews)

    final_matches = store.list_bank_income_matches(company_code="SL", period_yyyymm=period_yyyymm)
    counts = {
        "validated": 0,
        "suggested": 0,
        "pending_review": 0,
        "conflict": 0,
        "ignored": 0,
    }
    for match in final_matches:
        status = str(match.get("status") or "")
        if status in counts:
            counts[status] += 1

    return BankIncomeRunResult(
        company_code="SL",
        period_yyyymm=period_yyyymm,
        statement_dir=str(statement_dir),
        files_processed=len(caixabank_files),
        transactions_imported=len(current_transactions),
        validated_count=counts["validated"],
        suggested_count=counts["suggested"],
        pending_review_count=counts["pending_review"],
        conflict_count=counts["conflict"],
        ignored_count=counts["ignored"],
        review_count=len(store.list_bank_income_reviews(company_code="SL", period_yyyymm=period_yyyymm)),
    )


def _statement_dir(*, root: Path, company_code: str, period_yyyymm: str) -> Path:
    year = period_yyyymm[:4]
    folder = company_folder_name("ARTESTA STORE, S.L.") if company_code == "SL" else company_code
    return root / folder / year / period_yyyymm / "statements" / "bank"


def _parse_caixabank_statement(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _parse_caixabank_xls(path)
    if suffix == ".xlsx":
        return _parse_caixabank_xlsx(path)
    if suffix == ".csv":
        return _parse_caixabank_csv(path)
    raise ValueError(f"Unsupported statement format: {path}")


def _parse_caixabank_xls(path: Path) -> list[dict[str, Any]]:
    if xlrd is None:
        raise RuntimeError("xlrd is required to parse .xls Caixabank statements.")
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    raw_rows = [sheet.row_values(idx) for idx in range(sheet.nrows)]
    return _parse_caixabank_tabular_rows(
        raw_rows,
        parse_date=lambda value: _excel_date_xls(book, value),
    )


def _parse_caixabank_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        raw_rows = [list(raw) for raw in ws.iter_rows(values_only=True)]
        return _parse_caixabank_tabular_rows(raw_rows, parse_date=_excel_date_xlsx)
    finally:
        wb.close()


def _parse_caixabank_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        raw_rows = [list(raw) for raw in reader]
    return _parse_caixabank_tabular_rows(
        raw_rows,
        parse_date=lambda value: _parse_string_date(str(value or "").strip()),
    )


def _parse_caixabank_tabular_rows(
    raw_rows: list[list[Any]],
    *,
    parse_date,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    header_indexes: dict[str, int] | None = None
    for idx, raw in enumerate(raw_rows, start=1):
        values = list(raw)
        if header_indexes is None:
            detected = _detect_caixabank_header(values)
            if detected:
                header_indexes = detected
                continue
        if header_indexes is None and idx <= 2:
            continue
        if len(values) < 6:
            continue
        parsed = _row_to_transaction(
            concept=str(_value_at(values, header_indexes, "concept", 0) or "").strip(),
            booking_date=parse_date(_value_at(values, header_indexes, "booking_date", 1)),
            value_date=parse_date(_value_at(values, header_indexes, "value_date", 2)),
            detail=str(_value_at(values, header_indexes, "detail", 3) or "").strip(),
            amount=_to_decimal(_value_at(values, header_indexes, "amount", 4)),
            balance=_to_decimal(_value_at(values, header_indexes, "balance", 5)),
        )
        if parsed:
            rows.append(parsed)
    return rows


def _detect_caixabank_header(values: list[Any]) -> dict[str, int] | None:
    indexes: dict[str, int] = {}
    for idx, value in enumerate(values):
        normalized = _normalize_caixabank_header(value)
        for field, aliases in _CAIXABANK_HEADER_ALIASES.items():
            if normalized in aliases:
                indexes[field] = idx
                break
    if len(indexes) == len(_CAIXABANK_HEADER_ALIASES):
        return indexes
    return None


def _normalize_caixabank_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(char)
    )
    return re.sub(r"\s+", " ", text)


def _value_at(values: list[Any], header_indexes: dict[str, int] | None, field: str, fallback_index: int) -> Any:
    index = header_indexes[field] if header_indexes and field in header_indexes else fallback_index
    return values[index] if index < len(values) else None


def _row_to_transaction(
    *,
    concept: str,
    booking_date: date | None,
    value_date: date | None,
    detail: str,
    amount: Decimal | None,
    balance: Decimal | None,
) -> dict[str, Any] | None:
    if not concept and amount is None:
        return None
    if amount is None:
        return None
    return {
        "concept": concept,
        "booking_date": booking_date.isoformat() if booking_date else "",
        "value_date": value_date.isoformat() if value_date else "",
        "detail": detail,
        "amount": amount.quantize(_TWO_PLACES),
        "balance": balance.quantize(_TWO_PLACES) if balance is not None else None,
    }


def _classify_transaction(row: dict[str, Any]) -> tuple[str, bool, str]:
    amount = Decimal(str(row["amount"]))
    if amount <= _ZERO:
        return "non_income", True, "non_income"
    text = f"{row.get('concept', '')} {row.get('detail', '')}".strip()
    if _PAYPAL_PATTERN.search(text):
        return "paypal_income_ignored", True, "paypal_out_of_scope"
    if _SHOPIFY_PATTERN.search(text):
        return "shopify_payout_candidate", False, ""
    if _MANUAL_PATTERN.search(text):
        return "manual_transfer_candidate", False, ""
    return "unknown_income", False, ""


def _match_shopify_payout(
    *,
    tx: dict[str, Any],
    payouts: list[dict[str, Any]],
    used_payout_ids: set[str],
    current_period: str,
) -> tuple[dict[str, Any], dict[str, Any] | None, str | None]:
    tx_amount = Decimal(str(tx["amount"]))
    tx_date = _parse_iso_date(str(tx.get("booking_date") or ""))
    exact_candidates = [
        payout
        for payout in payouts
        if payout["target_id"] not in used_payout_ids
        and abs(Decimal(str(payout["net_amount"])) - tx_amount) <= _TWO_PLACES
    ]
    if len(exact_candidates) == 1:
        payout = exact_candidates[0]
        payout_date = _parse_iso_date(str(payout.get("payout_date") or ""))
        day_gap = abs((tx_date - payout_date).days) if tx_date and payout_date else 999
        outside_period = str(payout.get("payout_period") or "") != current_period
        if day_gap <= 7:
            status = "validated"
            reason_code = "shopify_outside_period" if outside_period else "shopify_matched"
            confidence = 98
        else:
            status = "suggested"
            reason_code = "shopify_outside_period" if outside_period else "shopify_amount_mismatch"
            confidence = 70
        return (
            _match_record(
                tx,
                status=status,
                reason_code=reason_code,
                match_type="shopify_payout",
                target_ids=[payout["target_id"]],
                confidence=confidence,
                outside_period=outside_period,
                notes=f"Matched Shopify payout {payout['target_id']}",
            ),
            None,
            payout["target_id"],
        )
    if len(exact_candidates) > 1:
        target_ids = [item["target_id"] for item in exact_candidates[:5]]
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="shopify_amount_mismatch",
            target_ids=target_ids,
            notes="Multiple Shopify payouts match the same bank amount.",
        )
        return (
            _match_record(
                tx,
                status="conflict",
                reason_code="shopify_amount_mismatch",
                match_type="shopify_payout",
                target_ids=target_ids,
                confidence=30,
            ),
            review,
            None,
        )
    near_candidates = [
        payout for payout in payouts
        if payout["target_id"] not in used_payout_ids
        and abs(Decimal(str(payout["net_amount"])) - tx_amount) <= Decimal("5.00")
    ]
    if near_candidates:
        payout = near_candidates[0]
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="shopify_amount_mismatch",
            target_ids=[payout["target_id"]],
            notes="Shopify-like income found but amount differs from expected payout.",
        )
        return (
            _match_record(
                tx,
                status="pending_review",
                reason_code="shopify_amount_mismatch",
                match_type="shopify_payout",
                target_ids=[payout["target_id"]],
                confidence=40,
            ),
            review,
            None,
        )
    review = _review_record(
        company_code="SL",
        period_yyyymm=current_period,
        bank_transaction_id=tx["id"],
        review_type="transaction",
        reason_code="shopify_missing_payout",
        notes="Stripe income without Shopify payout candidate.",
    )
    return (
        _match_record(
            tx,
            status="pending_review",
            reason_code="shopify_missing_payout",
            match_type="",
            target_ids=[],
            confidence=0,
        ),
        review,
        None,
    )


def _match_manual_income(
    *,
    tx: dict[str, Any],
    open_orders: list[dict[str, Any]],
    used_manual_order_ids: set[str],
    current_period: str,
) -> tuple[dict[str, Any], dict[str, Any] | None, list[str]]:
    tx_amount = Decimal(str(tx["amount"]))
    tx_date = _parse_iso_date(str(tx.get("booking_date") or ""))
    available_orders = [
        order for order in open_orders
        if order["order_name"] not in used_manual_order_ids
        and _is_within_manual_payment_window(tx_date=tx_date, order=order)
    ]
    exact_single = [
        order for order in available_orders
        if abs(Decimal(str(order["total"])) - tx_amount) <= _TWO_PLACES
    ]
    if len(exact_single) == 1:
        order = exact_single[0]
        outside_period = str(order.get("period_yyyymm") or "") != current_period
        return (
            _match_record(
                tx,
                status="validated",
                reason_code="manual_outside_period" if outside_period else "manual_matched",
                match_type="manual_order",
                target_ids=[order["order_name"]],
                confidence=95,
                outside_period=outside_period,
                notes=f"Matched manual order {order['order_name']}",
            ),
            None,
            [order["order_name"]],
        )
    if len(exact_single) > 1:
        preferred_order = _preferred_exact_manual_order(
            exact_single,
            tx_date=tx_date,
            current_period=current_period,
        )
        if preferred_order is not None:
            outside_period = str(preferred_order.get("period_yyyymm") or "") != current_period
            return (
                _match_record(
                    tx,
                    status="validated",
                    reason_code="manual_outside_period" if outside_period else "manual_matched",
                match_type="manual_order",
                target_ids=[preferred_order["order_name"]],
                confidence=80,
                outside_period=outside_period,
                notes=f"Matched manual order {preferred_order['order_name']} using closest-date heuristic.",
            ),
            None,
            [preferred_order["order_name"]],
            )
        target_ids = [order["order_name"] for order in exact_single[:5]]
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="manual_multiple_candidates",
            target_ids=target_ids,
            notes="More than one manual order matches the same incoming amount.",
        )
        return (
            _match_record(
                tx,
                status="conflict",
                reason_code="manual_multiple_candidates",
                match_type="manual_order",
                target_ids=target_ids,
                confidence=25,
            ),
            review,
            [],
        )

    combinations = _exact_order_combinations(available_orders, target_amount=tx_amount)
    if len(combinations) == 1:
        target_ids = [order["order_name"] for order in combinations[0]]
        outside_period = any(str(order.get("period_yyyymm") or "") != current_period for order in combinations[0])
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="manual_group_match_suggested",
            target_ids=target_ids,
            notes="One bank transfer likely covers multiple manual orders.",
        )
        return (
            _match_record(
                tx,
                status="suggested",
                reason_code="manual_group_match_suggested",
                match_type="manual_order_group",
                target_ids=target_ids,
                confidence=65,
                outside_period=outside_period,
            ),
            review,
            target_ids,
        )
    if len(combinations) > 1:
        target_ids = [order["order_name"] for order in combinations[0]]
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="manual_multiple_candidates",
            target_ids=target_ids,
            notes="Several manual order combinations match the same bank income.",
        )
        return (
            _match_record(
                tx,
                status="conflict",
                reason_code="manual_multiple_candidates",
                match_type="manual_order_group",
                target_ids=target_ids,
                confidence=20,
            ),
            review,
            [],
        )

    larger_orders = [
        order for order in available_orders
        if Decimal(str(order["total"])) > tx_amount
    ]
    if larger_orders:
        review = _review_record(
            company_code="SL",
            period_yyyymm=current_period,
            bank_transaction_id=tx["id"],
            review_type="transaction",
            reason_code="manual_partial_payment",
            notes="Incoming transfer may be a partial payment for an open manual order.",
        )
        return (
            _match_record(
                tx,
                status="pending_review",
                reason_code="manual_partial_payment",
                match_type="",
                target_ids=[],
                confidence=15,
            ),
            review,
            [],
        )

    review = _review_record(
        company_code="SL",
        period_yyyymm=current_period,
        bank_transaction_id=tx["id"],
        review_type="transaction",
        reason_code="manual_income_unmatched",
        notes="Incoming transfer did not match any open manual order.",
    )
    return (
        _match_record(
            tx,
            status="pending_review",
            reason_code="manual_income_unmatched",
            match_type="",
            target_ids=[],
            confidence=0,
        ),
        review,
        [],
    )


def _exact_order_combinations(open_orders: list[dict[str, Any]], *, target_amount: Decimal) -> list[list[dict[str, Any]]]:
    candidates = [
        order for order in open_orders
        if Decimal(str(order["total"])) <= target_amount
    ]
    candidates = sorted(candidates, key=lambda item: (str(item.get("order_date") or ""), item["order_name"]))[:30]
    matches: list[list[dict[str, Any]]] = []
    for size in (2, 3):
        for combo in itertools.combinations(candidates, size):
            combo_total = sum((Decimal(str(order["total"])) for order in combo), _ZERO)
            if abs(combo_total - target_amount) <= _TWO_PLACES:
                matches.append(list(combo))
    return matches


def _preferred_exact_manual_order(
    orders: list[dict[str, Any]],
    *,
    tx_date: date | None,
    current_period: str,
) -> dict[str, Any] | None:
    dated_orders = []
    for order in orders:
        order_date = _parse_iso_date(str(order.get("order_date") or ""))
        if not _is_within_manual_payment_window(tx_date=tx_date, order_date=order_date):
            continue
        delta_days = abs((tx_date - order_date).days) if tx_date and order_date else 9999
        dated_orders.append((order, order_date, str(order.get("period_yyyymm") or ""), delta_days))

    if not dated_orders:
        return None

    most_recent_period = max(period for _order, _order_date, period, _delta_days in dated_orders)
    period_orders = [
        (order, order_date, delta_days)
        for order, order_date, period, delta_days in dated_orders
        if period == most_recent_period
    ]
    period_orders.sort(key=lambda item: (item[2], item[1] or date.min, str(item[0].get("order_name") or "")))
    return period_orders[0][0] if period_orders else None


def _is_within_manual_payment_window(
    *,
    tx_date: date | None,
    order: dict[str, Any] | None = None,
    order_date: date | None = None,
) -> bool:
    if tx_date is None:
        return False
    if order_date is None and order is not None:
        order_date = _parse_iso_date(str(order.get("order_date") or ""))
    if order_date is None:
        return False
    return abs((tx_date - order_date).days) <= _MAX_MANUAL_PAYMENT_DELTA_DAYS


def _match_record(
    tx: dict[str, Any],
    *,
    status: str,
    reason_code: str,
    match_type: str,
    target_ids: list[str],
    confidence: int,
    outside_period: bool = False,
    notes: str = "",
    is_excluded: bool = False,
    exclusion_reason: str = "",
) -> dict[str, Any]:
    return {
        "id": _stable_id("bank-income-match", tx["id"]),
        "bank_transaction_id": tx["id"],
        "company_code": str(tx["company_code"]),
        "period_yyyymm": str(tx["period_yyyymm"]),
        "status": status,
        "reason_code": reason_code,
        "match_type": match_type,
        "target_ids": target_ids,
        "confidence": confidence,
        "outside_period": outside_period,
        "notes": notes,
        "is_excluded": is_excluded,
        "exclusion_reason": exclusion_reason,
        "source": "auto",
    }


def _review_record(
    *,
    company_code: str,
    period_yyyymm: str,
    bank_transaction_id: str,
    review_type: str,
    reason_code: str,
    target_ids: list[str] | None = None,
    notes: str = "",
) -> dict[str, Any]:
    key = "|".join([company_code, period_yyyymm, bank_transaction_id, review_type, reason_code, ",".join(target_ids or [])])
    return {
        "id": _stable_id("bank-income-review", key),
        "company_code": company_code,
        "period_yyyymm": period_yyyymm,
        "bank_transaction_id": bank_transaction_id,
        "review_type": review_type,
        "status": "open",
        "reason_code": reason_code,
        "target_ids": target_ids or [],
        "notes": notes,
        "source": "auto",
    }


def _transaction_fingerprint(row: dict[str, Any]) -> str:
    payload = "|".join([
        str(row.get("booking_date") or ""),
        str(row.get("value_date") or ""),
        _format_decimal(row["amount"]),
        _format_decimal(row["balance"]) if row["balance"] is not None else "",
        str(row.get("concept") or "").strip(),
        str(row.get("detail") or "").strip(),
    ])
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _stable_id(*parts: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, ":".join(parts)))


def _file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _to_decimal(value: object) -> Decimal | None:
    if value in (None, ""):
        return None
    if isinstance(value, Decimal):
        return value.quantize(_TWO_PLACES)
    if isinstance(value, (int, float)):
        return Decimal(str(value)).quantize(_TWO_PLACES)
    text = str(value).strip()
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(_TWO_PLACES)
    except Exception:
        return None


def _format_decimal(value: Decimal | str | int | float) -> str:
    return str(Decimal(str(value)).quantize(_TWO_PLACES))


def _excel_date_xls(book, value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        year, month, day, *_rest = xlrd.xldate_as_tuple(value, book.datemode)
        return date(year, month, day)
    return _parse_string_date(str(value))


def _excel_date_xlsx(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    return _parse_string_date(str(value))


def _parse_string_date(value: str) -> date | None:
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _parse_iso_date(value: str) -> date | None:
    return _parse_string_date(value)


def _previous_period(period_yyyymm: str) -> str:
    year = int(period_yyyymm[:4])
    month = int(period_yyyymm[4:])
    if month == 1:
        return f"{year - 1}12"
    return f"{year}{month - 1:02d}"
