from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from lector_facturas.bank_income_reconciliation import (
    _classify_transaction,
    _match_manual_income,
    _parse_caixabank_statement,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
VALIDATED_FILL = PatternFill("solid", fgColor="E2EFDA")
CONFLICT_FILL = PatternFill("solid", fgColor="FCE4D6")
PENDING_FILL = PatternFill("solid", fgColor="FFF2CC")

WHITE_BOLD = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
BOLD = Font(bold=True, name="Calibri", size=10)
NORMAL = Font(name="Calibri", size=10)
THIN_SIDE = Side(style="thin", color="D9D9D9")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)
MONEY_FMT = '#,##0.00;[Red](#,##0.00);-'


@dataclass(frozen=True)
class ManualIncomeReviewRow:
    booking_date: str
    period_yyyymm: str
    amount: Decimal
    concept: str
    detail: str
    status: str
    reason_code: str
    match_type: str
    target_ids: list[str]
    outside_period: bool


@dataclass(frozen=True)
class ManualIncomeReviewReport:
    company_code: str
    from_date: str
    to_date: str
    source_path: str
    rows: list[ManualIncomeReviewRow]


def collect_manual_income_review(
    store,
    *,
    company_code: str,
    statement_path: Path,
    from_date: str,
    to_date: str,
) -> ManualIncomeReviewReport:
    parsed_rows = _parse_caixabank_statement(statement_path)
    review_rows: list[ManualIncomeReviewRow] = []
    used_manual_order_ids: set[str] = set()

    candidate_rows = []
    for index, row in enumerate(parsed_rows, start=1):
        classification, _is_excluded, _exclusion_reason = _classify_transaction(row)
        booking_date = str(row.get("booking_date") or "")
        if classification != "manual_transfer_candidate":
            continue
        if booking_date < from_date or booking_date > to_date:
            continue
        candidate_rows.append(
            {
                "id": f"manual-income-{index}",
                "company_code": company_code,
                "period_yyyymm": booking_date[:7].replace("-", ""),
                "booking_date": booking_date,
                "value_date": str(row.get("value_date") or ""),
                "amount": str(row["amount"]),
                "concept": str(row.get("concept") or ""),
                "detail": str(row.get("detail") or ""),
                "classification": classification,
                "is_excluded": False,
                "exclusion_reason": "",
            }
        )
    candidate_rows.sort(key=lambda item: (item["booking_date"], item["amount"], item["detail"]))

    for tx in candidate_rows:
        open_orders = _list_manual_open_orders_without_consumed_matches(
            store,
            company_code=company_code,
            period_yyyymm=tx["period_yyyymm"],
        )
        match, _review, consumed_orders = _match_manual_income(
            tx=tx,
            open_orders=open_orders,
            used_manual_order_ids=used_manual_order_ids,
            current_period=tx["period_yyyymm"],
        )
        if consumed_orders:
            used_manual_order_ids.update(consumed_orders)
        review_rows.append(
            ManualIncomeReviewRow(
                booking_date=tx["booking_date"],
                period_yyyymm=tx["period_yyyymm"],
                amount=Decimal(str(tx["amount"])),
                concept=tx["concept"],
                detail=tx["detail"],
                status=str(match["status"]),
                reason_code=str(match["reason_code"]),
                match_type=str(match["match_type"]),
                target_ids=[str(item) for item in match.get("target_ids", [])],
                outside_period=bool(match.get("outside_period", False)),
            )
        )

    return ManualIncomeReviewReport(
        company_code=company_code,
        from_date=from_date,
        to_date=to_date,
        source_path=str(statement_path),
        rows=review_rows,
    )


def build_manual_income_review_workbook(report: ManualIncomeReviewReport) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    _build_summary_sheet(summary_ws, report)
    _build_detail_sheet(wb, "Validated", [row for row in report.rows if row.status == "validated"], VALIDATED_FILL)
    _build_detail_sheet(wb, "Conflicts", [row for row in report.rows if row.status == "conflict"], CONFLICT_FILL)
    _build_detail_sheet(wb, "Pending", [row for row in report.rows if row.status == "pending_review"], PENDING_FILL)
    _build_detail_sheet(wb, "All Rows", report.rows, None)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _list_manual_open_orders_without_consumed_matches(store, *, company_code: str, period_yyyymm: str) -> list[dict[str, Any]]:
    if getattr(store, "database_url", None):
        return store._list_manual_open_orders_db(company_code=company_code, period_yyyymm=period_yyyymm)
    return [
        dict(item)
        for item in store._read_bank_manual_orders()
        if str(item.get("company_code", "")) == company_code and str(item.get("period_yyyymm", "")) <= period_yyyymm
    ]


def _build_summary_sheet(ws, report: ManualIncomeReviewReport) -> None:
    widths = [14, 16, 12, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    _write_cell(ws, 1, 1, "B2B Manual Income Review", font=WHITE_BOLD, fill=HEADER_FILL, align=LEFT)
    ws.merge_cells("A1:D1")
    _write_cell(ws, 2, 1, f"Company: {report.company_code}")
    _write_cell(ws, 3, 1, f"Window: {report.from_date} to {report.to_date}")
    _write_cell(ws, 4, 1, f"Source: {report.source_path}")
    ws.merge_cells("A4:D4")

    row = 6
    _write_cell(ws, row, 1, "Status", font=BOLD, fill=SUBHEADER_FILL)
    _write_cell(ws, row, 2, "Count", font=BOLD, fill=SUBHEADER_FILL, align=RIGHT)
    row += 1
    counts = Counter(item.status for item in report.rows)
    for status in ("validated", "conflict", "pending_review"):
        _write_cell(ws, row, 1, status)
        _write_cell(ws, row, 2, counts.get(status, 0), align=RIGHT)
        row += 1

    row += 1
    _write_cell(ws, row, 1, "Month", font=BOLD, fill=SUBHEADER_FILL)
    _write_cell(ws, row, 2, "Validated", font=BOLD, fill=SUBHEADER_FILL, align=RIGHT)
    _write_cell(ws, row, 3, "Conflict", font=BOLD, fill=SUBHEADER_FILL, align=RIGHT)
    _write_cell(ws, row, 4, "Pending", font=BOLD, fill=SUBHEADER_FILL, align=RIGHT)
    row += 1
    monthly = {}
    for item in report.rows:
        monthly.setdefault(item.period_yyyymm, Counter())[item.status] += 1
    for period in sorted(monthly):
        _write_cell(ws, row, 1, period)
        _write_cell(ws, row, 2, monthly[period].get("validated", 0), align=RIGHT)
        _write_cell(ws, row, 3, monthly[period].get("conflict", 0), align=RIGHT)
        _write_cell(ws, row, 4, monthly[period].get("pending_review", 0), align=RIGHT)
        row += 1


def _build_detail_sheet(wb: Workbook, title: str, rows: list[ManualIncomeReviewRow], fill: PatternFill | None) -> None:
    ws = wb.create_sheet(title)
    headers = [
        "Date",
        "Period",
        "Amount",
        "Concept",
        "Detail",
        "Status",
        "Reason",
        "Match type",
        "Target IDs",
        "Outside period",
    ]
    widths = [12, 10, 12, 22, 34, 14, 24, 18, 40, 14]
    for idx, (header, width) in enumerate(zip(headers, widths, strict=True), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
        _write_cell(ws, 1, idx, header, font=WHITE_BOLD, fill=HEADER_FILL, align=CENTER)

    for row_idx, item in enumerate(rows, start=2):
        _write_cell(ws, row_idx, 1, item.booking_date, fill=fill)
        _write_cell(ws, row_idx, 2, item.period_yyyymm, fill=fill)
        _write_cell(ws, row_idx, 3, float(item.amount), fill=fill, align=RIGHT, number_format=MONEY_FMT)
        _write_cell(ws, row_idx, 4, item.concept, fill=fill)
        _write_cell(ws, row_idx, 5, item.detail, fill=fill)
        _write_cell(ws, row_idx, 6, item.status, fill=fill)
        _write_cell(ws, row_idx, 7, item.reason_code, fill=fill)
        _write_cell(ws, row_idx, 8, item.match_type, fill=fill)
        _write_cell(ws, row_idx, 9, ", ".join(item.target_ids), fill=fill)
        _write_cell(ws, row_idx, 10, "Yes" if item.outside_period else "", fill=fill, align=CENTER)


def _write_cell(
    ws,
    row: int,
    column: int,
    value: Any,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    align: Alignment = LEFT,
    number_format: str | None = None,
) -> None:
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = font or NORMAL
    cell.border = THIN_BORDER
    cell.alignment = align
    if fill:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format

